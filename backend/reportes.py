"""Exportar (Excel / PDF) e importar desde Excel."""
import calendar
import logging
from datetime import date
from io import BytesIO

import pandas as pd
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from alertas import correo_configurado, destinatarios, enviar_correo
from auth import usuario_actual
from database import ALMACEN_ID, conectar, get_db, transaccion
from servicios import aplicar_movimiento, auditar

log = logging.getLogger("inventario.reportes")

router = APIRouter(prefix="/api", dependencies=[Depends(usuario_actual)])

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 5000


def _descarga(buf: BytesIO, nombre: str, media: str):
    buf.seek(0)
    return StreamingResponse(buf, media_type=media, headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


def _excel_bytes(hojas: dict[str, pd.DataFrame]) -> BytesIO:
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for nombre, df in hojas.items():
            df.to_excel(writer, index=False, sheet_name=nombre)
            ws = writer.sheets[nombre]
            ws.freeze_panes = "A2"
            for i, col in enumerate(df.columns, start=1):
                largo = max([len(str(col))] + [len(str(x)) for x in df[col].head(500)])
                ws.column_dimensions[get_column_letter(i)].width = min(largo + 2, 50)
            for celda in ws[1]:
                celda.font = Font(bold=True)
    return buf


# ───────────────────────── Exportar ─────────────────────────
@router.get("/exportar/excel")
def exportar_excel(db=Depends(get_db)):
    df = pd.read_sql_query(
        """SELECT p.sku AS SKU, p.nombre AS Producto, p.categoria AS Categoría,
                  v.talla AS Talla, v.color AS Color,
                  COALESCE(v.stock_actual, 0) AS Stock, p.stock_minimo AS "Stock mínimo",
                  p.precio_costo AS Costo, p.precio_venta AS Venta,
                  COALESCE(v.stock_actual, 0) * p.precio_costo AS "Valor en stock"
           FROM productos p LEFT JOIN variantes v ON v.sku_producto = p.sku
           ORDER BY p.nombre COLLATE NOCASE, v.talla, v.color""",
        db,
    )
    return _descarga(_excel_bytes({"Inventario": df}), f"inventario_{date.today()}.xlsx", XLSX)


@router.get("/exportar/kardex")
def exportar_kardex(desde: str = "", hasta: str = "", db=Depends(get_db)):
    where, params = [], []
    if desde:
        where.append("date(k.fecha, 'localtime') >= ?")
        params.append(desde)
    if hasta:
        where.append("date(k.fecha, 'localtime') <= ?")
        params.append(hasta)
    cond = (" WHERE " + " AND ".join(where)) if where else ""
    df = pd.read_sql_query(
        f"""SELECT datetime(k.fecha, 'localtime') AS Fecha, k.sku_producto AS SKU, p.nombre AS Producto,
                   v.talla AS Talla, v.color AS Color, k.tipo_movimiento AS Tipo, k.cantidad AS Cantidad,
                   k.stock_anterior AS "Stock anterior", k.stock_resultante AS "Stock resultante",
                   k.referencia AS Referencia, k.usuario AS Usuario
            FROM kardex k
            LEFT JOIN productos p ON p.sku = k.sku_producto
            LEFT JOIN variantes v ON v.id = k.id_variante
            {cond} ORDER BY k.fecha DESC, k.id DESC""",
        db,
        params=params,
    )
    return _descarga(_excel_bytes({"Kardex": df}), f"kardex_{date.today()}.xlsx", XLSX)


def _generar_excel_rango_bytes(db, desde: str = "", hasta: str = "", q: str = "", categoria: str = "") -> BytesIO:
    where, params = [], []
    if desde:
        where.append("date(k.fecha, 'localtime') >= ?")
        params.append(desde)
    if hasta:
        where.append("date(k.fecha, 'localtime') <= ?")
        params.append(hasta)
    if q.strip():
        like = f"%{q.strip()}%"
        where.append("(k.sku_producto LIKE ? OR p.nombre LIKE ? OR k.referencia LIKE ?)")
        params += [like, like, like]
    if categoria.strip():
        where.append("p.categoria = ?")
        params.append(categoria.strip())

    cond = (" WHERE " + " AND ".join(where)) if where else ""

    df_kardex = pd.read_sql_query(
        f"""SELECT datetime(k.fecha, 'localtime') AS Fecha, k.sku_producto AS SKU, p.nombre AS Producto,
                   p.categoria AS Categoría, v.talla AS Talla, v.color AS Color,
                   k.tipo_movimiento AS Tipo, k.cantidad AS Cantidad,
                   k.stock_anterior AS "Stock anterior", k.stock_resultante AS "Stock resultante",
                   k.referencia AS Referencia, k.usuario AS Usuario
            FROM kardex k
            LEFT JOIN productos p ON p.sku = k.sku_producto
            LEFT JOIN variantes v ON v.id = k.id_variante
            {cond} ORDER BY k.fecha DESC, k.id DESC""",
        db,
        params=params,
    )

    df_resumen = pd.read_sql_query(
        f"""SELECT k.sku_producto AS SKU, p.nombre AS Producto, p.categoria AS Categoría,
                   SUM(CASE WHEN k.tipo_movimiento = 'ENTRADA' THEN k.cantidad ELSE 0 END) AS "Total Entradas",
                   SUM(CASE WHEN k.tipo_movimiento = 'SALIDA' THEN k.cantidad ELSE 0 END) AS "Total Salidas",
                   SUM(CASE WHEN k.tipo_movimiento = 'AJUSTE' THEN k.cantidad ELSE 0 END) AS "Ajustes Directos"
            FROM kardex k
            LEFT JOIN productos p ON p.sku = k.sku_producto
            {cond}
            GROUP BY k.sku_producto, p.nombre, p.categoria
            ORDER BY p.nombre COLLATE NOCASE""",
        db,
        params=params,
    )

    where_inv, params_inv = [], []
    if q.strip():
        like = f"%{q.strip()}%"
        where_inv.append("(p.sku LIKE ? OR p.nombre LIKE ?)")
        params_inv += [like, like]
    if categoria.strip():
        where_inv.append("p.categoria = ?")
        params_inv.append(categoria.strip())
    cond_inv = (" WHERE " + " AND ".join(where_inv)) if where_inv else ""

    df_inventario = pd.read_sql_query(
        f"""SELECT p.sku AS SKU, p.nombre AS Producto, p.categoria AS Categoría,
                  v.talla AS Talla, v.color AS Color,
                  COALESCE(v.stock_actual, 0) AS Stock, p.stock_minimo AS "Stock mínimo",
                  p.precio_costo AS Costo, p.precio_venta AS Venta,
                  COALESCE(v.stock_actual, 0) * p.precio_costo AS "Valor en stock"
           FROM productos p LEFT JOIN variantes v ON v.sku_producto = p.sku
           {cond_inv}
           ORDER BY p.nombre COLLATE NOCASE, v.talla, v.color""",
        db,
        params=params_inv,
    )

    hojas = {
        "Movimientos en Rango": df_kardex,
        "Resumen por Producto": df_resumen,
        "Inventario Filtrado": df_inventario,
    }
    return _excel_bytes(hojas)


@router.get("/exportar/excel-rango")
def exportar_excel_rango(desde: str = "", hasta: str = "", q: str = "", categoria: str = "", db=Depends(get_db)):
    buf = _generar_excel_rango_bytes(db, desde, hasta, q, categoria)
    nombre = f"reporte_inventario_{desde or 'inicio'}_a_{hasta or 'fin'}.xlsx"
    return _descarga(buf, nombre, XLSX)


@router.post("/reportes/enviar-mensual")
def enviar_reporte_mensual(
    anio: int = Query(..., ge=2000, le=2100),
    mes: int = Query(..., ge=1, le=12),
    destino: str = Query(""),
    db=Depends(get_db),
    usuario: str = Depends(usuario_actual),
):
    if not correo_configurado():
        raise HTTPException(400, "El correo no está configurado en el servidor. Configura las variables SMTP_* en backend/.env.")

    _, ultimo_dia = calendar.monthrange(anio, mes)
    desde = f"{anio:04d}-{mes:02d}-01"
    hasta = f"{anio:04d}-{mes:02d}-{ultimo_dia:02d}"

    buf = _generar_excel_rango_bytes(db, desde, hasta)
    nombre_archivo = f"reporte_mensual_{anio}_{mes:02d}.xlsx"

    asunto = f"📊 Reporte Mensual de Inventario - {mes:02d}/{anio}"
    cuerpo = (
        f"Reporte consolidado de inventario y movimientos correspondiente al período del {desde} al {hasta}.\n\n"
        f"Generado por: {usuario}\n"
        f"Fecha de emisión: {date.today():%d/%m/%Y}\n"
    )

    error = enviar_correo(
        asunto=asunto,
        texto=cuerpo,
        destino=destino or None,
        adjunto_bytes=buf.getvalue(),
        adjunto_nombre=nombre_archivo,
    )

    if error:
        raise HTTPException(502, f"Error al enviar el correo: {error}")

    auditar(db, usuario, f"Envió reporte mensual en Excel ({mes:02d}/{anio}) por correo")
    return {"enviado": True, "mensaje": f"Reporte del mes {mes:02d}/{anio} enviado con éxito por correo."}


def verificar_y_enviar_reporte_mensual_automatico():
    """Se ejecuta en segundo plano. Si es día 1 del mes y no se ha enviado el reporte del mes anterior, lo envía automáticamente."""
    hoy = date.today()
    if hoy.month == 1:
        mes_ant, anio_ant = 12, hoy.year - 1
    else:
        mes_ant, anio_ant = hoy.month - 1, hoy.year

    periodo = f"{anio_ant:04d}-{mes_ant:02d}"

    conn = conectar()
    try:
        cur = conn.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS reportes_enviados (
                periodo TEXT PRIMARY KEY,
                fecha_envio TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        ya_enviado = cur.execute("SELECT 1 FROM reportes_enviados WHERE periodo = ?", (periodo,)).fetchone()

        if hoy.day == 1 and not ya_enviado:
            if correo_configurado():
                _, ultimo_dia = calendar.monthrange(anio_ant, mes_ant)
                desde = f"{periodo}-01"
                hasta = f"{periodo}-{ultimo_dia:02d}"

                buf = _generar_excel_rango_bytes(conn, desde, hasta)
                asunto = f"📊 [AUTOMÁTICO] Reporte Mensual de Inventario - {mes_ant:02d}/{anio_ant}"
                cuerpo = (
                    f"Reporte mensual automatizado correspondiente al período {desde} al {hasta}.\n\n"
                    f"Este correo fue generado automáticamente por el Sistema de Inventario el primer día del mes.\n"
                )

                error = enviar_correo(
                    asunto=asunto,
                    texto=cuerpo,
                    adjunto_bytes=buf.getvalue(),
                    adjunto_nombre=f"reporte_mensual_{periodo}.xlsx",
                )

                if not error:
                    cur.execute("INSERT INTO reportes_enviados (periodo) VALUES (?)", (periodo,))
                    cur.execute("INSERT INTO auditoria (usuario, accion) VALUES (?, ?)", ("Sistema Automático", f"Envió reporte mensual automático ({periodo})"))
                    log.info("Reporte mensual automático %s enviado con éxito.", periodo)
                else:
                    log.warning("No se pudo enviar el reporte mensual automático %s: %s", periodo, error)
    except Exception as exc:
        log.exception("Error comprobando reporte mensual automático: %s", exc)
    finally:
        conn.close()


@router.get("/reportes/estado-automatico")
def estado_reporte_automatico(db=Depends(get_db)):
    hoy = date.today()
    if hoy.month == 1:
        mes_ant, anio_ant = 12, hoy.year - 1
    else:
        mes_ant, anio_ant = hoy.month - 1, hoy.year
    periodo_anterior = f"{anio_ant:04d}-{mes_ant:02d}"

    enviado = db.execute("SELECT fecha_envio FROM reportes_enviados WHERE periodo = ?", (periodo_anterior,)).fetchone()
    ultimos = db.execute("SELECT periodo, fecha_envio FROM reportes_enviados ORDER BY fecha_envio DESC LIMIT 5").fetchall()

    return {
        "programado": True,
        "frecuencia": "Día 1 de cada mes a las 00:00 UTC/Local",
        "correo_configurado": correo_configurado(),
        "destinatarios": destinatarios() if correo_configurado() else [],
        "periodo_anterior": periodo_anterior,
        "enviado_periodo_anterior": bool(enviado),
        "fecha_ultimo_envio": enviado["fecha_envio"] if enviado else None,
        "historial": [dict(u) for u in ultimos],
    }


@router.get("/exportar/pdf")
def exportar_pdf(db=Depends(get_db)):
    filas = db.execute(
        """SELECT p.sku, p.nombre, p.categoria, p.stock_minimo, p.precio_costo, p.precio_venta,
                  COALESCE(SUM(v.stock_actual), 0) AS stock
           FROM productos p LEFT JOIN variantes v ON v.sku_producto = p.sku
           GROUP BY p.sku ORDER BY p.nombre COLLATE NOCASE"""
    ).fetchall()

    estilos = getSampleStyleSheet()
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    datos = [["SKU", "Producto", "Categoría", "Stock", "Mín.", "Costo", "Venta"]]
    bajos = []
    for i, f in enumerate(filas, start=1):
        datos.append([
            f["sku"],
            Paragraph(f["nombre"], estilos["BodyText"]),
            f["categoria"] or "",
            str(f["stock"]),
            str(f["stock_minimo"]),
            f"{f['precio_costo'] or 0:,.2f}",
            f"{f['precio_venta'] or 0:,.2f}",
        ])
        if f["stock"] <= f["stock_minimo"]:
            bajos.append(i)

    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F1E2B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F7")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDE3EA")),
    ]
    for i in bajos:
        estilo.append(("TEXTCOLOR", (3, i), (3, i), colors.HexColor("#B91C1C")))
        estilo.append(("FONTNAME", (3, i), (3, i), "Helvetica-Bold"))

    tabla = Table(datos, repeatRows=1, colWidths=[2.6 * cm, 6 * cm, 3 * cm, 1.5 * cm, 1.3 * cm, 1.8 * cm, 1.8 * cm])
    tabla.setStyle(TableStyle(estilo))
    doc.build([
        Paragraph("Reporte de inventario", estilos["Title"]),
        Paragraph(f"Generado el {date.today():%d/%m/%Y}. En rojo: productos en o bajo su stock mínimo.", estilos["Normal"]),
        Spacer(1, 0.5 * cm),
        tabla,
    ])
    return _descarga(buf, f"inventario_{date.today()}.pdf", "application/pdf")


# ───────────────────────── Importar desde Excel ─────────────────────────
COLUMNAS = ["sku", "nombre", "categoria", "precio_costo", "precio_venta", "stock_minimo", "talla", "color", "stock"]
OBLIGATORIAS = ["sku", "nombre", "precio_costo", "precio_venta"]


@router.get("/importar/plantilla")
def descargar_plantilla():
    ejemplo = pd.DataFrame(
        [
            ["POL-001", "Polo básico algodón", "Polos", 12.5, 29.9, 5, "M", "Negro", 10],
            ["POL-001", "Polo básico algodón", "Polos", 12.5, 29.9, 5, "L", "Negro", 8],
        ],
        columns=COLUMNAS,
    )
    ayuda = pd.DataFrame({"Cómo llenar esta plantilla": [
        "Cada fila es UNA variante (talla + color) de un producto.",
        "Si un producto tiene varias variantes, repite sus datos (sku, nombre, precios) en cada fila.",
        "Obligatorias: sku, nombre, precio_costo y precio_venta.",
        "Opcionales: categoria, stock_minimo (por defecto 5), talla (por defecto Única), color (por defecto Único) y stock (por defecto 0).",
        "Si el SKU ya existe, sus datos no se modifican: solo se agregan las variantes nuevas.",
        "Borra las dos filas de ejemplo antes de subir el archivo.",
    ]})
    return _descarga(_excel_bytes({"Inventario": ejemplo, "Instrucciones": ayuda}), "plantilla_inventario.xlsx", XLSX)


def _numero(texto: str, nombre: str, entero: bool, errores: list[str], defecto=None):
    texto = str(texto).strip().replace(",", ".")
    if texto == "":
        if defecto is None:
            errores.append(f"{nombre} es obligatorio")
        return defecto
    try:
        valor = float(texto)
    except ValueError:
        errores.append(f"{nombre} no es un número")
        return defecto
    if valor < 0:
        errores.append(f"{nombre} no puede ser negativo")
        return defecto
    if entero:
        if not valor.is_integer():
            errores.append(f"{nombre} debe ser un número entero")
            return defecto
        return int(valor)
    return valor


def _analizar(db, contenido: bytes):
    """Valida el archivo completo sin tocar la base de datos."""
    try:
        df = pd.read_excel(BytesIO(contenido), dtype=str, sheet_name=0).fillna("")
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo. Sube un Excel .xlsx (puedes bajar la plantilla).")

    df.columns = [str(c).strip().lower() for c in df.columns]
    faltan = [c for c in OBLIGATORIAS if c not in df.columns]
    if faltan:
        raise HTTPException(400, "Faltan columnas obligatorias: " + ", ".join(faltan) + ". Usa la plantilla.")
    if len(df) > MAX_FILAS:
        raise HTTPException(400, f"El archivo tiene más de {MAX_FILAS} filas. Divídelo en partes.")

    productos_db = {r["sku"] for r in db.execute("SELECT sku FROM productos")}
    variantes_db = {
        (r["sku_producto"], (r["talla"] or "").lower(), (r["color"] or "").lower())
        for r in db.execute("SELECT sku_producto, talla, color FROM variantes")
    }

    errores, validas = [], []
    productos_archivo: dict[str, tuple[tuple, int]] = {}
    vistas: dict[tuple, int] = {}

    def celda(fila, col):
        return str(fila.get(col, "")).strip()

    for idx, fila in df.iterrows():
        n = idx + 2  # fila real en Excel (1 = encabezado)
        if all(str(v).strip() == "" for v in fila):
            continue
        errs: list[str] = []
        sku = celda(fila, "sku").upper()
        nombre = celda(fila, "nombre")
        if not sku:
            errs.append("sku es obligatorio")
        if not nombre:
            errs.append("nombre es obligatorio")
        costo = _numero(celda(fila, "precio_costo"), "precio_costo", False, errs)
        venta = _numero(celda(fila, "precio_venta"), "precio_venta", False, errs)
        minimo = _numero(celda(fila, "stock_minimo"), "stock_minimo", True, errs, defecto=5)
        stock = _numero(celda(fila, "stock"), "stock", True, errs, defecto=0)
        talla = celda(fila, "talla") or "Única"
        color = celda(fila, "color") or "Único"
        categoria = celda(fila, "categoria")

        if not errs:
            datos = (nombre, categoria, costo, venta, minimo)
            clave = (sku, talla.lower(), color.lower())
            if clave in vistas:
                errs.append(f"la variante {talla}/{color} de {sku} está repetida (ya aparece en la fila {vistas[clave]})")
            if clave in variantes_db:
                errs.append(f"la variante {talla}/{color} de {sku} ya existe en el inventario")
            if sku in productos_archivo and productos_archivo[sku][0] != datos:
                errs.append(f"los datos de {sku} no coinciden con los de la fila {productos_archivo[sku][1]}")
        if errs:
            errores.append({"fila": n, "errores": errs})
            continue

        vistas[clave] = n
        productos_archivo.setdefault(sku, (datos, n))
        validas.append({"fila": n, "sku": sku, "talla": talla, "color": color, "stock": stock})

    nuevos = [s for s in productos_archivo if s not in productos_db]
    resumen = {
        "filas": len(validas) + len(errores),
        "productos_nuevos": len(nuevos),
        "productos_existentes": len(productos_archivo) - len(nuevos),
        "variantes_nuevas": len(validas),
        "unidades": sum(v["stock"] for v in validas),
    }
    return {"validas": validas, "productos": productos_archivo, "productos_db": productos_db, "errores": errores, "resumen": resumen}


@router.post("/importar/excel")
async def importar_excel(
    archivo: UploadFile,
    confirmar: bool = Query(False),
    db=Depends(get_db),
    usuario: str = Depends(usuario_actual),
):
    if not (archivo.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "El archivo debe ser un Excel .xlsx.")
    contenido = await archivo.read()
    if len(contenido) > MAX_BYTES:
        raise HTTPException(413, "El archivo pesa más de 5 MB.")

    r = _analizar(db, contenido)
    respuesta = {
        "valido": not r["errores"] and bool(r["validas"]),
        "aplicado": False,
        "resumen": r["resumen"],
        "errores": r["errores"][:100],
        "total_errores": len(r["errores"]),
    }
    if not confirmar:
        return respuesta

    if r["errores"]:
        raise HTTPException(400, "El archivo tiene errores. Corrígelos y vuelve a subirlo; no se importó nada.")
    if not r["validas"]:
        raise HTTPException(400, "El archivo no tiene filas para importar.")

    with transaccion(db):
        for sku, (datos, _) in r["productos"].items():
            if sku in r["productos_db"]:
                continue
            nombre, categoria, costo, venta, minimo = datos
            db.execute(
                "INSERT INTO productos (sku, nombre, categoria, precio_costo, precio_venta, stock_minimo) VALUES (?, ?, ?, ?, ?, ?)",
                (sku, nombre, categoria, costo, venta, minimo),
            )
        for v in r["validas"]:
            cur = db.execute(
                "INSERT INTO variantes (sku_producto, talla, color, stock_actual, id_sucursal) VALUES (?, ?, ?, 0, ?)",
                (v["sku"], v["talla"], v["color"], ALMACEN_ID),
            )
            if v["stock"] > 0:
                aplicar_movimiento(db, cur.lastrowid, "ENTRADA", v["stock"], "Importación Excel", usuario)
        auditar(
            db, usuario,
            f"Importó Excel: {r['resumen']['productos_nuevos']} productos nuevos, {r['resumen']['variantes_nuevas']} variantes",
        )
    respuesta["aplicado"] = True
    return respuesta